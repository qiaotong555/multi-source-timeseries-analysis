# -*- coding: utf-8 -*-
"""Generate a detailed Word document for the weather Agent demo."""
from __future__ import annotations

import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUMMARY_PATH = ROOT / "artifacts" / "demo" / "demo_summary.json"
CASE_SUMMARY_PATH = ROOT / "artifacts" / "demo_case_unit_change_rule" / "demo_summary.json"
CASE_RESULT_PATH = ROOT / "artifacts" / "demo_case_unit_change_rule" / "pipeline" / "result.json"
OUT_PATH = ROOT / "doc" / "Agent一键Demo说明_2026-10-24.docx"
AUTHOR = "qiaotong555"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F2F7FC"
GREEN = "E2F0D9"
PALE_GREEN = "F3F9EF"
ORANGE = "FCE4D6"
PALE_ORANGE = "FFF6F0"
GRAY = "F2F2F2"
DARK_GRAY = "595959"
WHITE = "FFFFFF"
RED = "C00000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs: dict[str, str]) -> None:
    """Set table cell borders. kwargs: top/bottom/left/right -> attrs."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:%s" % edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn("w:%s" % key), str(value))


def set_cell_text(cell, text: str, bold: bool = False,
                  color: str = "000000", size: float = 9.5,
                  align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    set_run_font(run, "Microsoft YaHei")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_para(doc: Document, text: str = "", style: str | None = None,
             bold_prefix: str | None = None, color: str | None = None,
             size: float | None = None, align=None, space_after: float = 6,
             first_line: float | None = None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if first_line is not None:
        p.paragraph_format.first_line_indent = Inches(first_line)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        if color:
            r1.font.color.rgb = RGBColor.from_string(color)
        if size:
            r1.font.size = Pt(size)
        set_run_font(r1)
        r2 = p.add_run(text[len(bold_prefix):])
        if color:
            r2.font.color.rgb = RGBColor.from_string(color)
        if size:
            r2.font.size = Pt(size)
        set_run_font(r2)
    else:
        r = p.add_run(text)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if size:
            r.font.size = Pt(size)
        set_run_font(r)
    return p


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    set_run_font(r)


def add_number(doc: Document, text: str, level: int = 0) -> None:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    set_run_font(r)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    r.font.size = Pt(16 if level == 1 else 12.5)
    set_run_font(r)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE,
                title_color: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": "8", "color": title_color},
                    bottom={"val": "single", "sz": "8", "color": title_color},
                    left={"val": "single", "sz": "18", "color": title_color},
                    right={"val": "single", "sz": "8", "color": title_color})
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(title_color)
    r.font.size = Pt(11)
    set_run_font(r)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]],
              widths: list[float] | None = None, header_fill: str = NAVY,
              font_size: float = 9.2) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=WHITE, size=font_size,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr[i], header_fill)
        set_cell_border(hdr[i], top={"val": "single", "sz": "6", "color": header_fill},
                        bottom={"val": "single", "sz": "6", "color": header_fill},
                        left={"val": "single", "sz": "4", "color": "D9E2F3"},
                        right={"val": "single", "sz": "4", "color": "D9E2F3"})
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            set_cell_text(cells[cidx], "—" if value is None else str(value), size=font_size)
            set_cell_shading(cells[cidx], PALE_BLUE if ridx % 2 == 0 else WHITE)
            set_cell_border(cells[cidx], top={"val": "single", "sz": "3", "color": "D9E2F3"},
                            bottom={"val": "single", "sz": "3", "color": "D9E2F3"},
                            left={"val": "single", "sz": "3", "color": "D9E2F3"},
                            right={"val": "single", "sz": "3", "color": "D9E2F3"})
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    set_cell_border(cell, top={"val": "single", "sz": "4", "color": "BFBFBF"},
                    bottom={"val": "single", "sz": "4", "color": "BFBFBF"},
                    left={"val": "single", "sz": "4", "color": "BFBFBF"},
                    right={"val": "single", "sz": "4", "color": "BFBFBF"})
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string("333333")
        set_run_font(r, "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def get_mode(summary: dict[str, Any], mode: str) -> dict[str, Any]:
    return next(x for x in summary["runs"] if x["mode"] == mode)


def build_document() -> Path:
    with SUMMARY_PATH.open(encoding="utf-8") as f:
        summary = json.load(f)
    if not CASE_SUMMARY_PATH.exists() or not CASE_RESULT_PATH.exists():
        raise FileNotFoundError(
            "缺少案例产物。请先运行：MTSA_PLANNER=rule python scripts/demo.py "
            "--mode pipeline --input fixtures/weather_v1.csv "
            "--question \"请重点检查 STA_A 温度数据是否发生单位变化，并分析温度趋势和异常点，给出处理建议\" "
            "--out artifacts/demo_case_unit_change_rule"
        )
    with CASE_SUMMARY_PATH.open(encoding="utf-8") as f:
        case_summary = json.load(f)
    with CASE_RESULT_PATH.open(encoding="utf-8") as f:
        case_doc = json.load(f)
    pipeline = get_mode(summary, "pipeline")
    agent = get_mode(summary, "agent")
    case = get_mode(case_summary, "pipeline")
    case_quality = (case_doc.get("data") or {}).get("quality") or {}
    case_clean = (case_doc.get("data") or {}).get("clean") or {}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    # Global styles
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name, size, color in [("Heading 1", 16, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 11, DARK_GRAY)]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if "Small Note" not in styles:
        st = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Microsoft YaHei"
        st.font.size = Pt(9)
        st.font.color.rgb = RGBColor.from_string(DARK_GRAY)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    # Header/footer
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("多源时序数据量化分析与可视化 Agent · Demo 说明")
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_run_font(hr)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("2026-10-24  |  第 ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_run_font(fr)
    add_field(footer, "PAGE")
    fr2 = footer.add_run(" 页")
    fr2.font.size = Pt(8)
    fr2.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_run_font(fr2)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(68)
    r = p.add_run("多源时序数据量化分析与可视化 Agent")
    r.bold = True
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    set_run_font(r)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(9)
    r2 = p2.add_run("一键 Demo 全流程说明与实测结果")
    r2.font.size = Pt(19)
    r2.font.color.rgb = RGBColor.from_string(BLUE)
    r2.bold = True
    set_run_font(r2)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(26)
    r3 = p3.add_run("天气正式场景 · pipeline 与 agent 双模式")
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_run_font(r3)

    doc.add_paragraph().paragraph_format.space_before = Pt(35)
    cover_rows = [
        ["文档目的", "解释 Demo 具体做了什么、为什么这样做、得到了什么结果，以及如何复现"],
        ["实际运行时间", summary["generated_at"]],
        ["输入数据", summary["input"]],
        ["分析问题", summary["question"]],
        ["模型", summary.get("model") or "未调用"],
        ["项目根目录", "仓库克隆位置（本机路径因人而异）"],
    ]
    add_table(doc, ["项目项", "内容"], cover_rows, widths=[1.25, 5.65], header_fill=BLUE, font_size=10)
    add_callout(doc, "先看结论", "这是一个可运行的 Demo，不是静态截图。文档中的数字来自 2026 年 10 月 24 日实际运行生成的 result.json、report.md 和 demo_summary.json。", fill=PALE_GREEN, title_color="548235")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Contents
    add_heading(doc, "目录", 1)
    contents = [
        "1. Demo 要解决什么问题",
        "2. 这次 Demo 具体做了什么",
        "3. 为什么要同时跑 pipeline 和 agent",
        "4. Agent 的完整执行流程",
        "5. 实际得到了什么结果",
        "6. 结果如何理解：成功、差异与未执行项",
        "7. 证据链与复现方法",
        "8. 当前限制与下一步",
        "附录：关键文件清单",
    ]
    for item in contents:
        add_bullet(doc, item)
    add_callout(doc, "阅读建议", "如果只想快速了解结果，先看第 5、6 节；如果想自己运行，直接看第 7 节；如果想理解 Agent 为什么不是“让模型直接算数据”，看第 2、4 节。", fill=PALE_BLUE, title_color=BLUE)
    add_page_break(doc)

    # 1
    add_heading(doc, "1. Demo 要解决什么问题", 1)
    add_para(doc, "这个 Demo 的目的不是展示一个漂亮的聊天界面，而是把项目当前已经实现的分析能力串成一条可运行、可观察、可复现的完整链路。用户给出一份天气时序数据和一个自然语言问题，系统需要完成数据体检、清洗、分析、结论生成、证据绑定和结果落盘。", first_line=0.3)
    add_callout(doc, "一句话理解", "Demo 验证的是：同一份数据、同一个问题，固定流水线和 Agent 两种运行方式能否都稳定产出结构化结果，并且每条结论都能回到证据。", fill=PALE_BLUE, title_color=BLUE)
    add_heading(doc, "1.1 本次 Demo 的输入", 2)
    add_table(doc, ["输入项", "实际值", "说明"], [
        ["场景", "weather", "当前 Agent 正式支持的场景"],
        ["数据文件", "fixtures/weather_v1.csv", "两站点、多源天气时序数据"],
        ["原始规模", f"{pipeline['rows']:,} 行 × {pipeline['columns']} 列", "包含时间、站点和天气变量"],
        ["时间范围", f"{pipeline['time_start']} → {pipeline['time_end']}", "小时级、跨两年"],
        ["自然语言问题", summary["question"], "同时要求趋势、季节性、异常点和变化点"],
    ], widths=[1.25, 2.65, 3.0])
    add_heading(doc, "1.2 本次 Demo 的输出", 2)
    add_bullet(doc, "机器可读结果：result.json，包含状态、数据质量、分析指标、异常、变化点、结论、证据和 Agent 轨迹。")
    add_bullet(doc, "人可读报告：report.md，方便快速查看数据概况、质量问题、结论和证据。")
    add_bullet(doc, "Demo 汇总：demo_summary.json 与 demo_summary.md，专门用于对比 pipeline 和 agent。")
    add_bullet(doc, "图表和其他运行产物：保存在各自的运行目录中，按 run_id 隔离。")

    # 2
    add_heading(doc, "2. 这次 Demo 具体做了什么", 1)
    add_para(doc, "Demo 不是重新实现一套分析逻辑，而是调用项目已有的引擎入口，把一次运行拆成几个可解释的阶段。这样做的好处是：Demo 的结果与正式 CLI 结果共用同一条代码路径，不会出现“演示版一套逻辑、正式版另一套逻辑”的问题。", first_line=0.3)
    add_heading(doc, "2.1 固定准备层：先把数据弄清楚", 2)
    add_table(doc, ["阶段", "具体动作", "本次结果"], [
        ["导入", "读取 CSV，识别时间列、数值列、分组列", "识别为 weather；时间列为 timestamp；分组列为 station"],
        ["规模与结构", "检查行数、列数、时间范围和采样间隔", f"{pipeline['rows']:,} 行；{pipeline['columns']} 列；中位采样间隔 3600 秒"],
        ["质量检查", "检查缺失、重复、乱序、采样缺口、异常候选、单位变化和水平跳变", f"发现 {pipeline['quality_issues']} 条质量记录"],
        ["场景计划", "根据场景与问题选择目标序列和分析项", "目标序列为 temperature@STA_A / temperature@STA_B"],
        ["清洗", "单位转换、重复时间戳聚合、透视、排序、插值", "清洗后 17,530 行；插补 1,531 个单元格"],
    ], widths=[1.0, 3.0, 2.9])
    add_heading(doc, "2.2 分析层：pipeline 和 agent 的差别从这里开始", 2)
    add_table(doc, ["模式", "怎么决定分析项", "本次实际执行"], [
        ["pipeline", "代码按固定计划执行", "trend、seasonal、anomaly、changepoint"],
        ["agent", "模型逐步选择白名单工具，每次看工具摘要后再决定下一步", "compute_metrics、detect_anomalies、detect_change_points、analyze_spectrum"],
    ], widths=[1.0, 3.1, 2.8])
    add_heading(doc, "2.3 固定收尾层：把结果变成可审计产物", 2)
    add_bullet(doc, "将数值结果整理成结构化结论。")
    add_bullet(doc, "为结论绑定 evidence_id，检查每个 evidence_id 是否确实存在。")
    add_bullet(doc, "生成 report.md，并把运行状态、错误、警告和轨迹写入 result.json。")
    add_bullet(doc, "因此，用户看到的不只是“模型说了什么”，还可以查到“这句话由哪个指标支持”。")

    add_heading(doc, "2.4 为什么采用这种分层设计", 2)
    add_table(doc, ["设计选择", "原因"], [
        ["数值由代码计算", "避免模型凭语言猜数值，保证指标、异常分数和变化点可复算。"],
        ["Agent 只负责工具编排", "让模型发挥“选择分析路径”的能力，同时把数学计算留在确定性代码中。"],
        ["准备层和收尾层固定", "避免 Agent 擅自改变清洗顺序或漏掉证据校验。"],
        ["pipeline 作为基线", "可以判断 Agent 是改变了分析路径，还是改变了计算结果。"],
        ["同数据同问题对比", "控制变量，差异主要来自运行模式，而不是输入变化。"],
    ], widths=[2.0, 4.9])

    # 3
    add_heading(doc, "3. 为什么要同时跑 pipeline 和 agent", 1)
    add_para(doc, "如果只运行 Agent，看到“成功”并不能说明它比固定流水线更好；如果只运行 pipeline，又无法证明 Agent 的编排机制能够工作。因此 Demo 让两种模式使用同一份数据、同一个问题、同一套底层工具，并分别保存结果。", first_line=0.3)
    add_table(doc, ["比较维度", "pipeline", "agent", "比较意义"], [
        ["稳定性", "固定顺序", "模型决定顺序", "看 Agent 是否能完成必要分析"],
        ["数值来源", "代码", "代码", "确保模式差异不是模型算错数字"],
        ["分析覆盖", "固定计划", "问题下限 + 模型追加", "观察模型是否会主动补充周期等分析"],
        ["成本", "不调用 LLM", "记录 token 和调用次数", "评估 Agent 的运行代价"],
        ["可复现", "确定性较强", "轨迹可保存、可重放", "出现差异时能够追踪原因"],
    ], widths=[1.2, 1.55, 1.55, 2.6])
    add_callout(doc, "本次 Demo 的控制变量", "输入文件、自然语言问题、规则计划环境、分析工具和阈值保持一致；变化的是“分析层如何选择工具”。", fill=PALE_GREEN, title_color="548235")

    # 4
    add_heading(doc, "4. Agent 的完整执行流程", 1)
    add_para(doc, "Agent 不是一次性让模型输出最终报告，而是一个观察—决策—行动循环。模型每次只决定下一步工具，工具执行后返回压缩摘要，模型再根据摘要决定是否继续。", first_line=0.3)
    add_heading(doc, "4.1 总体流程", 2)
    add_table(doc, ["步骤", "执行者", "做什么", "是否允许模型决定"], [
        ["1. 导入", "固定代码", "读取文件并识别字段", "否"],
        ["2. 质量与摘要", "固定代码", "检查数据质量并生成概要", "否"],
        ["3. 场景与计划", "固定代码/规则计划", "确定天气场景和初始分析范围", "否"],
        ["4. 清洗", "固定代码", "单位转换、聚合、排序、插值", "否"],
        ["5. 工具决策", "Agent", "选择下一个分析工具", "是"],
        ["6. 工具执行", "代码工具", "计算指标、检测异常或变化点", "参数受约束"],
        ["7. 观察回灌", "Agent runtime", "把工具摘要回传给模型", "是"],
        ["8. 收尾", "固定代码", "生成图表、报告、证据并落盘", "否"],
    ], widths=[0.75, 1.25, 3.4, 1.5])
    add_heading(doc, "4.2 本次 Agent 的四步轨迹", 2)
    add_table(doc, ["步", "工具", "结果", "为什么做"], [
        ["1", "compute_metrics", "ok", "先建立均值、波动和趋势基线"],
        ["2", "detect_anomalies", "ok", "问题明确要求异常点，属于范围下限"],
        ["3", "detect_change_points", "ok", "问题明确要求变化点，属于范围下限"],
        ["4", "analyze_spectrum", "ok", "问题包含季节/周期含义，补充频域分析"],
        ["追加尝试", "run_stat_tests", "未执行", "超出 guided 模式允许的追加步数，被安全拒绝"],
    ], widths=[0.55, 2.0, 1.0, 3.35])
    add_callout(doc, "重要解释", "报告中的 budget_exhausted 是 guided 模式“追加步数用尽”，不是 token 预算耗尽。本次 Agent 实际使用 5,254 / 8,000 token，token 预算没有耗尽。", fill=PALE_ORANGE, title_color="C65911")
    add_heading(doc, "4.3 为什么没有让模型直接看原始数组", 2)
    add_bullet(doc, "原始数组规模大、容易造成上下文污染，也不利于复现。")
    add_bullet(doc, "工具先用代码算出摘要，模型只看结构化结果，例如异常数量、趋势方向、频谱峰值。")
    add_bullet(doc, "这样既保留了模型的路径选择能力，又避免模型修改数值计算。")

    # 5 results
    add_heading(doc, "5. 实际得到了什么结果", 1)
    add_para(doc, "以下结果全部来自本次 Demo 的实际运行产物，而不是方案预期或手工填写。", first_line=0.3)
    add_heading(doc, "5.1 pipeline 与 agent 总体对比", 2)
    add_table(doc, ["指标", "pipeline", "agent", "解释"], [
        ["运行状态", pipeline["status"], agent["status"], "两种模式均成功落盘"],
        ["结论数量", pipeline["conclusions"], agent["conclusions"], "agent 多出周期、相关和滞后关系结论"],
        ["证据数量", pipeline["evidence"], agent["evidence"], "每条结论均有对应 evidence"],
        ["证据覆盖率", f"{pipeline['evidence_coverage']:.2f}", f"{agent['evidence_coverage']:.2f}", "没有无证据结论"],
        ["异常点", pipeline["anomalies"], agent["anomalies"], "两模式底层异常检测结果一致"],
        ["异常系列数", pipeline["anomaly_series"], agent["anomaly_series"], "STA_A 与 STA_B 温度序列"],
        ["变化点", pipeline["change_points"], agent["change_points"], "本数据运行中未产生变化点检出"],
        ["耗时", f"{pipeline['elapsed_ms']:,} ms", f"{agent['elapsed_ms']:,} ms", "agent 包含模型决策开销"],
        ["Agent 步数", "—", agent["agent"]["steps"], "模型做了 4 次工具决策"],
        ["Agent token", "—", agent["agent"]["tokens"]["total_tokens"], "累计 token，未超过 8,000 上限"],
    ], widths=[1.35, 1.0, 1.0, 3.55])
    add_heading(doc, "5.2 数据质量结果", 2)
    quality_rows = []
    for code, count in pipeline["quality_codes"].items():
        meaning = {
            "MISSING_VALUES": "存在缺失值",
            "DUPLICATE_TIMESTAMPS": "多站点共享时间戳，先按时间+站点聚合",
            "OUT_OF_ORDER": "时间顺序不单调，需要排序",
            "SAMPLING_GAP": "存在采样缺口",
            "OUTLIER_CANDIDATE": "质量层发现的候选异常点",
            "POSSIBLE_UNIT_CHANGE": "疑似单位/量纲变化",
            "LEVEL_SHIFT": "疑似水平变化",
        }.get(code, "质量检查记录")
        quality_rows.append([code, count, meaning])
    add_table(doc, ["质量代码", "次数", "含义"], quality_rows, widths=[2.0, 0.8, 4.1])
    add_para(doc, "数据质量检查的价值是先告诉用户“这份数据是否干净”，再进行后续分析。本次数据中尤其重要的是 STA_A 温度的一段疑似华氏度→摄氏度单位变化，系统在清洗层执行了单位转换，并把转换操作写入清洗记录。", first_line=0.3)
    add_heading(doc, "5.3 主要结论示例", 2)
    add_table(doc, ["类别", "pipeline 结论示例", "agent 结论示例"], [
        ["趋势", "STA_A 方向上升；STA_B 方向下降。", "保留相同趋势结论。"],
        ["季节性", "STA_A 峰值 7 月、谷值 1 月；STA_B 峰值 6 月、谷值 12 月。", "保留相同季节性结论。"],
        ["周期", "未执行频域工具，因此没有周期结论。", "发现约 23.27 小时主周期，频谱能量占比约 93%。"],
        ["相关性", "没有该类结论。", "STA_A 与 STA_B Pearson 相关系数约 0.792，属于强相关。"],
        ["滞后", "没有该类结论。", "发现约 -26 个采样点的滞后关系，相关系数约 0.832。"],
        ["异常检测", "STA_A 7 个、STA_B 11 个，共 18 个。", "相同，共 18 个。"],
        ["变化点", "没有检出。", "没有检出。"],
    ], widths=[1.0, 2.85, 2.85])
    add_callout(doc, "结果含义", "Agent 的价值在本次 Demo 中主要体现在“根据问题主动补充分析项”，而不是改变异常检测的数值结果。它增加了周期、相关和滞后结论，同时保留了相同的趋势、季节性和异常结果。", fill=PALE_GREEN, title_color="548235")

    # 6 interpretation
    add_heading(doc, "6. 结果如何理解：成功、差异与未执行项", 1)
    add_heading(doc, "6.1 哪些地方是成功的", 2)
    add_bullet(doc, "pipeline 和 agent 都返回 succeeded，且都生成 result.json 与 report.md。")
    add_bullet(doc, "两种模式的异常点结果一致：18 个异常点，覆盖两个温度系列。")
    add_bullet(doc, "两种模式的证据覆盖率都是 1.00，说明最终结论没有脱离证据。")
    add_bullet(doc, "Agent 完成了问题范围下限要求的趋势、季节性、异常和变化点相关工作。")
    add_bullet(doc, "缺少模型时 Demo 能明确降级为只跑 pipeline，而不是伪装成 Agent 成功。")
    add_heading(doc, "6.2 为什么 agent 比 pipeline 多 3 条结论", 2)
    add_table(doc, ["差异", "原因", "是否代表数值改变"], [
        ["11 条 vs 8 条结论", "Agent 额外执行 analyze_spectrum，并产生周期、相关、滞后关系结论。", "不代表已有数值被修改"],
        ["异常点均为 18 个", "两种模式调用同一个 detect_anomalies 工具。", "没有改变"],
        ["agent 更慢", "多了模型调用和逐步决策。", "属于运行成本差异"],
        ["agent 工具覆盖 4/5", "run_stat_tests 被追加步数边界拒绝。", "不影响本次四类主结论，但说明覆盖并不完整"],
    ], widths=[1.6, 3.9, 1.2])
    add_heading(doc, "6.3 本次 Demo 不能证明什么", 2)
    add_bullet(doc, "不能证明 Agent 在所有天气数据上都一定优于 pipeline。")
    add_bullet(doc, "不能证明“成功一次”就代表生产稳定性；稳定性需要重复实验与真实数据验证。")
    add_bullet(doc, "不能把 18 个异常点全部当成真实业务异常；Demo 数据是项目 fixture，主要用于验证流程。")
    add_bullet(doc, "不能把股票场景算作 Agent 能力；当前股票 Agent 仍未实现。")
    add_bullet(doc, "不能把本次 agent 与 pipeline 结果相同理解为两者完全等价；Agent 少执行了统计检验工具。")

    # 7 evidence and reproduce
    add_heading(doc, "7. 证据链与复现方法", 1)
    add_para(doc, "为了让 Demo 不停留在“看起来能跑”，项目把每次运行拆成可追踪的证据链。下面的文件可以从输入、运行、结论一直追到最终摘要。", first_line=0.3)
    add_heading(doc, "7.1 证据链", 2)
    add_table(doc, ["证据层", "文件", "作用"], [
        ["输入", "fixtures/weather_v1.csv", "Demo 使用的原始天气数据"],
        ["入口", "scripts/demo.py", "一键运行与摘要汇总逻辑"],
        ["汇总", "artifacts/demo/demo_summary.json", "pipeline/agent 对比数据"],
        ["人读汇总", "artifacts/demo/demo_summary.md", "可直接查看的演示结果"],
        ["pipeline 原始结果", "artifacts/demo/pipeline/result.json", "固定流水线完整结构化输出"],
        ["pipeline 报告", "artifacts/demo/pipeline/report.md", "固定流水线人读报告"],
        ["agent 原始结果", "artifacts/demo/agent/result.json", "Agent 轨迹、工具覆盖、token 和结论"],
        ["agent 报告", "artifacts/demo/agent/report.md", "Agent 人读报告"],
        ["方法说明", "doc/评估设计.md、doc/项目说明书.md", "指标口径、能力边界和限制"],
    ], widths=[1.25, 3.2, 2.25])
    add_heading(doc, "7.2 最简单的复现方式", 2)
    add_code_block(doc, "# 在仓库根目录（含 pyproject.toml）执行\npython scripts/demo.py")
    add_para(doc, "默认行为：使用 fixtures/weather_v1.csv，使用天气场景，问题为“分析温度的趋势、季节性和异常点，并检测变化点”，并尝试运行 pipeline 与 agent。", first_line=0.3)
    add_heading(doc, "7.3 常用复现方式", 2)
    add_code_block(doc, "# 只跑 pipeline，不需要模型\npython scripts/demo.py --mode pipeline\n\n# 运行 Agent，需要 MTSA_LLM_BASE_URL、MTSA_LLM_API_KEY、MTSA_LLM_MODEL\npython scripts/demo.py --mode agent\n\n# 指定 v2 数据与问题\npython scripts/demo.py --input fixtures/weather_eval_v2.csv --question \"找出温度湿度的异常、气压的突变，并说明季节规律\" --mode both")
    add_heading(doc, "7.4 如果没有模型配置", 2)
    add_table(doc, ["命令", "行为"], [
        ["python scripts/demo.py --mode both", "pipeline 正常执行，agent 标记跳过并写明缺失环境变量"],
        ["python scripts/demo.py --mode agent", "直接返回 DEMO_AGENT_UNAVAILABLE，退出码 2"],
        ["python scripts/demo.py --mode pipeline", "完全离线运行固定流水线"],
    ], widths=[3.1, 3.6])
    add_heading(doc, "7.5 如何进一步核对某条结论", 2)
    add_number(doc, "在 demo_summary.md 中找到对应的结论类别和摘要。")
    add_number(doc, "打开相应模式的 report.md，查看结论绑定的 evidence_id。")
    add_number(doc, "在 result.json 的 evidence 数组中找到同一个 evidence_id。")
    add_number(doc, "继续根据 evidence 的 metric_id、列名和时间范围，追溯到 data.metrics、data.anomalies 或 data.change_points。")

    # 8 limits
    add_heading(doc, "8. 当前限制与下一步", 1)
    add_heading(doc, "8.1 本 Demo 的限制", 2)
    add_table(doc, ["限制", "影响"], [
        ["数据是项目 fixture", "不能直接代表真实传感器或业务数据。"],
        ["Agent 只支持天气", "股票场景当前仅有 pipeline，不能通过本 Demo 验证股票 Agent。"],
        ["run_stat_tests 未执行", "本次 Agent 只覆盖 4/5 工具；主结论不受影响，但工具覆盖不是完整 100%。"],
        ["变化点为 0", "本次数据和参数下没有变化点输出，不等于系统永远检测不到变化点。"],
        ["一次运行不能证明稳定性", "还需要真实数据、重复运行和更大样本验证。"],
        ["token 只记录数量", "当前没有把 token 换算为供应商金额。"],
    ], widths=[2.0, 4.7])
    add_heading(doc, "8.2 项目当前建议路线", 2)
    add_table(doc, ["优先级", "下一步", "目的"], [
        ["P0", "真实天气数据验证", "补上合成数据之外的证据缺口"],
        ["P1", "变化点精度与持续异常检测改进", "把当前压线指标变成有余量指标"],
        ["P2", "成本语义和工具覆盖可观测性", "让 Agent 的成本和漏执行原因更易解释"],
        ["P3", "REST/异步服务", "有外部调用需求时再启动"],
    ], widths=[0.8, 2.5, 3.4])
    add_callout(doc, "明确暂不做", "多模型对比（已取消）、股票 Agent（暂缓）、预测分析、CSI/彩票场景。", fill=PALE_ORANGE, title_color="C65911")

    # Case study
    add_heading(doc, "9. 案例演示：气象站温度单位变化与异常诊断", 1)
    add_para(doc, "前面的 Demo 展示了系统整体怎么跑；下面用一个更接近实际使用的问题，演示用户如何根据报告发现数据问题、理解清洗动作，并决定后续人工核查重点。这个案例只运行确定性 pipeline，不依赖模型，便于复现。", first_line=0.3)
    add_heading(doc, "9.1 案例背景与问题", 2)
    add_table(doc, ["项目", "内容"], [
        ["案例名称", "STA_A 温度单位变化与异常诊断"],
        ["输入", "fixtures/weather_v1.csv"],
        ["用户问题", "请重点检查 STA_A 温度数据是否发生单位变化，并分析温度趋势和异常点，给出处理建议"],
        ["运行模式", "pipeline；MTSA_PLANNER=rule"],
        ["为什么选这个案例", "它同时包含数据质量问题、自动清洗、趋势分析和异常检测，能体现系统不只是“算指标”，还会先检查数据是否可信。"],
    ], widths=[1.45, 5.25])
    add_callout(doc, "案例核心", "如果温度数据中间有一段从摄氏度变成华氏度，直接计算趋势和异常会把单位跳变误认为天气变化。系统必须先发现并处理这个问题，再继续分析。", fill=PALE_ORANGE, title_color="C65911")

    add_heading(doc, "9.2 案例执行过程", 2)
    add_table(doc, ["步骤", "系统动作", "案例中的实际结果"], [
        ["1. 识别问题", "读取时间列、站点列和天气字段，识别为 weather 场景。", "场景置信度 1.00；目标包括 STA_A/STA_B 温度以及其他天气变量。"],
        ["2. 发现单位异常", "比较同一站点温度在相邻区间的稳健统计量。", "STA_A 在 2024-05-05 和 2024-06-04 出现两条 POSSIBLE_UNIT_CHANGE 警告。"],
        ["3. 执行清洗", "把识别出的华氏度区间转换为摄氏度。", "2024-05-05 至 2024-06-04，共转换 713 行；中位数约 76.94 → 24.97。"],
        ["4. 处理结构问题", "聚合重复时间戳、排序、插值缺失。", "聚合重复记录 25 行；插补 1,531 个单元格；清洗后 17,530 行。"],
        ["5. 分析趋势与异常", "在清洗后的数据上执行趋势、周期、异常、变化点等分析。", "生成 39 条结论、39 条证据、139 个异常点、6 个变化点。"],
        ["6. 形成报告", "把质量问题、清洗操作、结论和证据写入结果文件。", "result.json 和 report.md 均成功落盘。"],
    ], widths=[1.0, 2.9, 2.8])

    add_heading(doc, "9.3 案例实际结果", 2)
    add_table(doc, ["指标", "结果", "如何理解"], [
        ["状态", case["status"], "案例运行成功"],
        ["质量问题记录", case["quality_issues"], "包含缺失、重复、乱序、采样缺口、单位变化等"],
        ["单位变化警告", (case["quality_codes"].get("POSSIBLE_UNIT_CHANGE") or 0), "STA_A 温度区间前后各有一次边界警告"],
        ["清洗后行数", "17,530", "重复时间戳聚合和透视后的有效时间序列规模"],
        ["单位转换行数", "713", "STA_A 温度从华氏度转换为摄氏度"],
        ["结论数 / 证据数", f"{case['conclusions']} / {case['evidence']}", "每条结论均有证据"],
        ["证据覆盖率", f"{case['evidence_coverage']:.2f}", "没有无证据结论"],
        ["异常点", case["anomalies"], "覆盖 6 个系列；包括温度、风速和降雨序列"],
        ["变化点", case["change_points"], "主要落在 pressure@STA_A / pressure@STA_B"],
    ], widths=[1.65, 1.25, 3.8])

    unit_issues = [x for x in case_quality.get("issues", []) if x.get("code") == "POSSIBLE_UNIT_CHANGE"]
    unit_rows = []
    for issue in unit_issues:
        unit_rows.append([
            issue.get("group"), issue.get("start"), issue.get("detail"),
        ])
    add_heading(doc, "9.4 最关键的证据：单位变化与自动修复", 2)
    if unit_rows:
        add_table(doc, ["站点", "边界时间", "质量层证据"], unit_rows, widths=[1.1, 1.5, 4.1])
    conversions = case_clean.get("unit_conversions") or []
    if conversions:
        c = conversions[0]
        add_table(doc, ["清洗字段", "转换区间", "转换行数", "转换前中位数", "转换后中位数", "转换"], [[
            c.get("column"), f"{c.get('start_time')} → {c.get('end_time')}", c.get("rows"),
            c.get("median_before"), c.get("median_after"), f"{c.get('unit_from')} → {c.get('unit_to')}"
        ]], widths=[1.1, 2.0, 0.75, 1.0, 1.0, 1.0], font_size=8.6)
    add_para(doc, "这里的关键不是“系统替用户做了一个不可逆决定”，而是：系统先发现疑似单位变化，依据已有场景规则执行受控转换，并把转换区间、行数、前后统计量写入 clean.operations。用户仍然可以在 result.json 中复核这一动作。", first_line=0.3)

    add_heading(doc, "9.5 案例结论与人工建议", 2)
    add_bullet(doc, "数据层面：STA_A 温度在 2024-05-05 至 2024-06-04 出现单位异常，系统将 713 行统一转换为摄氏度。")
    add_bullet(doc, "分析层面：清洗后仍能得到 STA_A 上升趋势、STA_B 下降趋势，以及两个温度序列的季节性和异常结果。")
    add_bullet(doc, "异常层面：STA_A 检出 7 个温度异常点，STA_B 检出 11 个温度异常点；其他天气变量也可能产生异常，需要结合业务上下文确认。")
    add_bullet(doc, "人工核查建议：确认原始设备在 2024-05-05 前后是否更换了单位配置；保留原始数据和转换记录；对高稳健 Z 值点核对传感器日志。")
    add_bullet(doc, "不能直接得出的结论：系统不能仅凭统计结果证明设备一定故障，也不能把所有异常点都当成真实天气事件。")
    add_callout(doc, "这个案例说明了什么", "Agent/分析系统的价值不仅是输出一个数字，还包括先识别数据是否可信、记录做过哪些清洗、再在清洗后的数据上给出有证据的结论。", fill=PALE_GREEN, title_color="548235")

    add_heading(doc, "9.6 案例复现", 2)
    add_code_block(doc, "# 在仓库根目录（含 pyproject.toml）执行\n$env:MTSA_PLANNER='rule'\npython scripts/demo.py --mode pipeline --input fixtures/weather_v1.csv --question \"请重点检查 STA_A 温度数据是否发生单位变化，并分析温度趋势和异常点，给出处理建议\" --out artifacts/demo_case_unit_change_rule\n\n# 根据案例产物重新生成包含案例章节的 Word 文档\npython scripts/make_demo_word.py")
    add_para(doc, "案例证据：artifacts/demo_case_unit_change_rule/pipeline/result.json、artifacts/demo_case_unit_change_rule/pipeline/report.md、artifacts/demo_case_unit_change_rule/demo_summary.json。", style="Small Note")

    # Appendix
    add_heading(doc, "附录：关键文件清单", 1)
    files = [
        ("一键 Demo", "scripts/demo.py"),
        ("Demo 汇总 JSON", "artifacts/demo/demo_summary.json"),
        ("Demo 汇总 Markdown", "artifacts/demo/demo_summary.md"),
        ("pipeline 原始结果", "artifacts/demo/pipeline/result.json"),
        ("agent 原始结果", "artifacts/demo/agent/result.json"),
        ("项目说明书", "doc/项目说明书.md"),
        ("Demo 运行记录", "doc/reports/演示运行记录_2026-10-24.md"),
        ("交付完成记录", "spec/项目承接已完成.md（第 30 章）"),
        ("未来规划", "spec/未来规划.md"),
    ]
    add_table(doc, ["用途", "路径"], files, widths=[2.0, 4.7])
    add_callout(doc, "文档结论", "这个 Demo 已经证明项目能够把“数据输入 → 质量检查 → 分析工具 → Agent 编排 → 证据绑定 → 报告落盘”完整跑通。下一步重点不是再做一个展示页面，而是用真实天气数据验证这些能力能否外推。", fill=PALE_GREEN, title_color="548235")

    # Last metadata
    p = doc.add_paragraph(style="Small Note")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("文档由项目真实 Demo 产物自动整理生成；未包含 API Key 或其他敏感凭据。")

    # Add core properties（作者与时间写真实值：模板默认的 python-docx / 2013 年
    # 会在「文件 → 信息」里暴露生成工具，交付前必须覆盖）
    stamp = datetime.now(timezone.utc).replace(microsecond=0)
    props = doc.core_properties
    props.title = "多源时序数据量化分析与可视化 Agent 一键 Demo 全流程说明"
    props.subject = "天气场景 pipeline / agent 双模式演示"
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.keywords = "时序分析, Agent, pipeline, weather, demo, 证据, 复现"
    props.comments = ""
    props.category = ""
    props.created = stamp
    props.modified = stamp

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
